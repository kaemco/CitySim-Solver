"""
Generate a PowerPoint presentation of the Gaussian Plume Dispersion Model
equations implemented in CitySim's pollutant.cpp.
"""

import io, os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── colour palette ────────────────────────────────────────────────────────────
BG      = RGBColor(0x1A, 0x1A, 0x2E)   # deep navy
ACCENT  = RGBColor(0x16, 0x21, 0x3E)   # card background
BLUE    = RGBColor(0x53, 0xB8, 0xFF)   # highlight / title
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
YELLOW  = RGBColor(0xFF, 0xD7, 0x00)
GRAY    = RGBColor(0xB0, 0xB8, 0xC8)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

# ── helpers ───────────────────────────────────────────────────────────────────

def new_prs():
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs

def blank_slide(prs):
    layout = prs.slide_layouts[6]   # completely blank
    return prs.slides.add_slide(layout)

def fill_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, l, t, w, h, color):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_textbox(slide, text, l, t, w, h,
                size=18, bold=False, color=WHITE,
                align=PP_ALIGN.LEFT, italic=False):
    txb = slide.shapes.add_textbox(l, t, w, h)
    tf  = txb.text_frame
    tf.word_wrap = True
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txb

def render_latex(latex, fontsize=18, color="white", bg_color=None,
                 dpi=180, pad=0.12):
    """Render a LaTeX string to a PNG bytes buffer via matplotlib."""
    fig, ax = plt.subplots(figsize=(0.01, 0.01))
    if bg_color:
        fig.patch.set_facecolor(bg_color)
        ax.set_facecolor(bg_color)
    else:
        fig.patch.set_alpha(0)
        ax.set_alpha(0)
    ax.axis("off")
    t = ax.text(0.5, 0.5, f"${latex}$",
                fontsize=fontsize, color=color,
                ha="center", va="center",
                transform=ax.transAxes)
    # auto-size figure
    fig.canvas.draw()
    bbox = t.get_window_extent(renderer=fig.canvas.get_renderer())
    w_in = bbox.width  / dpi + 2*pad
    h_in = bbox.height / dpi + 2*pad
    fig.set_size_inches(w_in, h_in)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi,
                bbox_inches="tight",
                transparent=(bg_color is None),
                facecolor=fig.get_facecolor())
    plt.close(fig)
    buf.seek(0)
    return buf

def add_latex(slide, latex, cx, cy, max_w=Inches(11),
              fontsize=20, color="white"):
    """Render LaTeX and insert centred at (cx, cy)."""
    buf = render_latex(latex, fontsize=fontsize, color=color)
    from PIL import Image as PILImage
    img = PILImage.open(buf)
    pw, ph = img.size
    buf.seek(0)
    aspect = ph / pw
    w = min(max_w, Inches(pw / 180))
    h = w * aspect
    l = cx - w // 2
    t = cy - h // 2
    slide.shapes.add_picture(buf, l, t, w, h)

def section_bar(slide, title):
    """Top coloured bar with slide title."""
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.85), ACCENT)
    add_rect(slide, 0, 0, Inches(0.06), Inches(0.85), BLUE)
    add_textbox(slide, title,
                Inches(0.18), Inches(0.1),
                Inches(12.5), Inches(0.65),
                size=24, bold=True, color=BLUE)

def bullet(slide, text, left, top, width, size=15, color=GRAY, indent=False):
    prefix = "    •  " if indent else "•  "
    add_textbox(slide, prefix + text, left, top, width, Inches(0.4),
                size=size, color=color)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE BUILDERS
# ═══════════════════════════════════════════════════════════════════════════════

def slide_title(prs):
    s = blank_slide(prs)
    fill_bg(s)
    # large accent band
    add_rect(s, 0, Inches(2.6), SLIDE_W, Inches(2.4), ACCENT)
    add_rect(s, 0, Inches(2.6), Inches(0.12), Inches(2.4), BLUE)

    add_textbox(s, "Gaussian Plume Dispersion Model",
                Inches(0.25), Inches(2.75), Inches(12.5), Inches(1.0),
                size=36, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_textbox(s, "Physical Equations & Coefficients",
                Inches(0.25), Inches(3.65), Inches(12.5), Inches(0.55),
                size=22, bold=False, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(s, "CitySim · pollutant.cpp",
                Inches(0.25), Inches(4.3), Inches(12.5), Inches(0.4),
                size=14, italic=True, color=GRAY, align=PP_ALIGN.CENTER)
    # decorative bottom stripe
    add_rect(s, 0, Inches(7.1), SLIDE_W, Inches(0.4), BLUE)


def slide_coordinate(prs):
    s = blank_slide(prs)
    fill_bg(s)
    section_bar(s, "1 · Coordinate Rotation into Wind Frame")

    add_textbox(s, "Meteorological convention:  φ  is the direction FROM which the wind blows",
                Inches(0.4), Inches(1.0), Inches(12.5), Inches(0.38),
                size=14, color=GRAY, italic=True)

    eqs = [
        (r"\hat{w}_x = -\sin\phi \;,\quad \hat{w}_y = -\cos\phi",
         "Wind unit vector components (world frame)", Inches(8.0)),
        (r"d_x = x_r - x_e \;,\quad d_y = y_r - y_e",
         "Emittor → receptor vector components", Inches(8.0)),
        (r"x_{\parallel} = d_x(-\sin\phi) + d_y(-\cos\phi)",
         "Downwind distance  [m]", Inches(8.0)),
        (r"y_{\perp} = -d_x(-\cos\phi) + d_y(-\sin\phi)",
         "Crosswind offset  [m]", Inches(8.0)),
    ]

    tops = [Inches(1.55), Inches(2.5), Inches(3.5), Inches(4.5)]
    for (latex, label, max_w), top in zip(eqs, tops):
        add_latex(s, latex, SLIDE_W//2, top, max_w=max_w, fontsize=22)
        add_textbox(s, label,
                    Inches(0.3), top + Inches(0.42), Inches(12.5), Inches(0.35),
                    size=13, color=GRAY, align=PP_ALIGN.CENTER, italic=True)

    add_rect(s, Inches(0.3), Inches(5.6), Inches(12.7), Inches(0.01), BLUE)
    add_textbox(s, "If  x ≤ 0  →  receptor is upwind  →  C = 0.   Otherwise  x = max(x, x_min)",
                Inches(0.3), Inches(5.65), Inches(12.5), Inches(0.4),
                size=14, color=YELLOW, align=PP_ALIGN.CENTER)

    # rotation diagram
    add_textbox(s, "North",
                Inches(10.5), Inches(6.2), Inches(1.0), Inches(0.35),
                size=12, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(s, "↑  y-world",
                Inches(10.5), Inches(6.5), Inches(1.5), Inches(0.35),
                size=11, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(s, "→  x-world",
                Inches(10.5), Inches(6.85), Inches(1.5), Inches(0.35),
                size=11, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(s, "↗  downwind (x∥)",
                Inches(9.0), Inches(6.5), Inches(2.5), Inches(0.35),
                size=11, color=BLUE, align=PP_ALIGN.LEFT)


def slide_windspeed(prs):
    s = blank_slide(prs)
    fill_bg(s)
    section_bar(s, "2 · Wind Speed — Logarithmic Profile")

    add_latex(s,
        r"u(h) = \max\!\left(\; u_{ref} \cdot \frac{\ln(h/z_0)}{\ln(h_{ref}/z_0)},\; u_{min} \;\right)",
        SLIDE_W//2, Inches(2.3), max_w=Inches(11), fontsize=26)

    add_latex(s,
        r"h = \max\!\left(z_e + \Delta z_{plume},\; 1.01\cdot z_0\right)",
        SLIDE_W//2, Inches(3.6), max_w=Inches(9), fontsize=22)

    params = [
        ("u_ref",  "Climate wind speed at reference height"),
        ("h_ref",  "Reference height  (default 10 m)"),
        ("z₀",     "Aerodynamic roughness length  (m)"),
        ("u_min",  "Minimum wind speed floor  (default 0.5 m/s)"),
        ("z_e",    "Emittor elevation  (m)"),
        ("Δz_plume", "User-prescribed plume rise  (m)"),
    ]
    left  = Inches(1.2)
    top   = Inches(4.5)
    for sym, desc in params:
        add_textbox(s, f"  {sym}", left, top, Inches(1.8), Inches(0.33),
                    size=14, bold=True, color=BLUE)
        add_textbox(s, desc, left + Inches(1.8), top, Inches(9.5), Inches(0.33),
                    size=14, color=GRAY)
        top += Inches(0.38)


def slide_stability(prs):
    s = blank_slide(prs)
    fill_bg(s)
    section_bar(s, "3 · Pasquill Stability Class  (AUTO mode)")

    add_textbox(s, "Determined from global horizontal irradiance  I_gh  and wind speed  u",
                Inches(0.4), Inches(1.0), Inches(12.5), Inches(0.4),
                size=15, color=GRAY, italic=True)

    # table data
    headers = ["Condition", "u < 2 m/s", "2 ≤ u < 3", "3 ≤ u < 5", "u ≥ 5"]
    rows = [
        ["I_gh > 600 W/m²  (strong sun)",  "A", "B", "C", "D"],
        ["300 < I_gh ≤ 600  (moderate sun)", "B", "C", "D", "D"],
        ["10 < I_gh ≤ 300  (weak sun)",    "C", "D", "D", "D"],
        ["I_gh ≤ 10, cloud > 0.5  (overcast)", "D", "D", "D", "D"],
        ["I_gh ≤ 10, cloud ≤ 0.5  (clear night)", "F", "E", "D", "D"],
    ]

    col_w = [Inches(4.2), Inches(1.7), Inches(1.7), Inches(1.7), Inches(1.7)]
    col_x = [Inches(0.3)]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)

    row_h = Inches(0.52)
    top0  = Inches(1.55)

    # header row
    add_rect(s, Inches(0.3), top0, sum(col_w), row_h, BLUE)
    for j, (x, w, h) in enumerate(zip(col_x, col_w, headers)):
        add_textbox(s, h, x + Inches(0.05), top0 + Inches(0.08),
                    w - Inches(0.1), row_h,
                    size=13, bold=True, color=RGBColor(0,0,0), align=PP_ALIGN.CENTER)

    stability_colors = {"A": RGBColor(0xFF,0x6B,0x6B),
                        "B": RGBColor(0xFF,0xA5,0x00),
                        "C": RGBColor(0xFF,0xD7,0x00),
                        "D": RGBColor(0x53,0xB8,0xFF),
                        "E": RGBColor(0xA0,0xFF,0xA0),
                        "F": RGBColor(0xD0,0xD0,0xFF)}

    for i, row in enumerate(rows):
        top = top0 + (i+1)*row_h
        bg  = ACCENT if i % 2 == 0 else RGBColor(0x22,0x22,0x3A)
        add_rect(s, Inches(0.3), top, sum(col_w), row_h, bg)
        for j, (x, w, cell) in enumerate(zip(col_x, col_w, row)):
            color = stability_colors.get(cell, WHITE) if j > 0 else GRAY
            add_textbox(s, cell, x + Inches(0.05), top + Inches(0.08),
                        w - Inches(0.1), row_h,
                        size=13, bold=(j > 0), color=color,
                        align=PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT)

    add_textbox(s, "A = very unstable   B = unstable   C = slightly unstable   "
                   "D = neutral   E = slightly stable   F = stable",
                Inches(0.3), Inches(6.95), Inches(12.5), Inches(0.4),
                size=12, color=GRAY, italic=True, align=PP_ALIGN.CENTER)


def slide_sigma(prs):
    s = blank_slide(prs)
    fill_bg(s)
    section_bar(s, "4 · Pasquill–Gifford Dispersion Coefficients  σ_y, σ_z")

    add_textbox(s, "x = downwind distance (m)     f_y, f_z = optional user scaling factors",
                Inches(0.4), Inches(1.0), Inches(12.5), Inches(0.35),
                size=14, color=GRAY, italic=True)

    # table
    headers = ["Class", "σ_y  (m)", "σ_z  (m)"]
    rows = [
        ["A", "0.22 x (1 + 0.0001x)^{-1/2}", "0.20 x"],
        ["B", "0.16 x (1 + 0.0001x)^{-1/2}", "0.12 x"],
        ["C", "0.11 x (1 + 0.0001x)^{-1/2}", "0.08 x (1 + 0.0002x)^{-1/2}"],
        ["D", "0.08 x (1 + 0.0001x)^{-1/2}", "0.06 x (1 + 0.0015x)^{-1/2}"],
        ["E", "0.06 x (1 + 0.0001x)^{-1/2}", "0.03 x (1 + 0.0003x)^{-1}"],
        ["F", "0.04 x (1 + 0.0001x)^{-1/2}", "0.016 x (1 + 0.0003x)^{-1}"],
    ]
    col_w = [Inches(0.9), Inches(5.4), Inches(5.4)]
    col_x = [Inches(1.0)]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)

    row_h = Inches(0.55)
    top0  = Inches(1.45)

    add_rect(s, Inches(1.0), top0, sum(col_w), row_h, BLUE)
    for x, w, h in zip(col_x, col_w, headers):
        add_textbox(s, h, x + Inches(0.05), top0 + Inches(0.08),
                    w - Inches(0.1), row_h,
                    size=14, bold=True, color=RGBColor(0,0,0), align=PP_ALIGN.CENTER)

    stability_colors = {"A": RGBColor(0xFF,0x6B,0x6B),
                        "B": RGBColor(0xFF,0xA5,0x00),
                        "C": RGBColor(0xFF,0xD7,0x00),
                        "D": RGBColor(0x53,0xB8,0xFF),
                        "E": RGBColor(0xA0,0xFF,0xA0),
                        "F": RGBColor(0xD0,0xD0,0xFF)}

    for i, row in enumerate(rows):
        top = top0 + (i+1)*row_h
        bg  = ACCENT if i % 2 == 0 else RGBColor(0x22,0x22,0x3A)
        add_rect(s, Inches(1.0), top, sum(col_w), row_h, bg)
        for j, (x, w, cell) in enumerate(zip(col_x, col_w, row)):
            color = stability_colors.get(cell, GRAY) if j == 0 else WHITE
            add_textbox(s, cell, x + Inches(0.05), top + Inches(0.1),
                        w - Inches(0.1), row_h,
                        size=13, bold=(j == 0), color=color,
                        align=PP_ALIGN.CENTER if j == 0 else PP_ALIGN.LEFT)

    add_latex(s,
        r"\sigma_y^{\prime} = f_y \cdot \sigma_y \qquad \sigma_z^{\prime} = f_z \cdot \sigma_z",
        SLIDE_W//2, Inches(6.85), max_w=Inches(8), fontsize=18)


def slide_concentration(prs):
    s = blank_slide(prs)
    fill_bg(s)
    section_bar(s, "5 · Gaussian Plume Concentration")

    add_latex(s,
        r"C(x,y,z) = \frac{Q}{2\pi\,u\,\sigma_y\sigma_z}"
        r"\cdot\exp\!\left(-\frac{y^2}{2\sigma_y^2}\right)"
        r"\cdot T_z",
        SLIDE_W//2, Inches(2.0), max_w=Inches(12), fontsize=26)

    add_textbox(s, "Vertical term  T_z  —  with ground reflection (image source at –H):",
                Inches(0.5), Inches(3.0), Inches(12.0), Inches(0.4),
                size=15, color=GRAY)

    add_latex(s,
        r"T_z = \exp\!\left(-\frac{(z-H)^2}{2\sigma_z^2}\right)"
        r"+ \exp\!\left(-\frac{(z+H)^2}{2\sigma_z^2}\right)",
        SLIDE_W//2, Inches(3.85), max_w=Inches(11), fontsize=24)

    add_latex(s,
        r"H = z_e + \Delta z_{plume}",
        SLIDE_W//2, Inches(4.9), max_w=Inches(7), fontsize=20)

    params = [
        ("Q",  "Emission rate  [kg/s]"),
        ("u",  "Wind speed at effective height  [m/s]"),
        ("H",  "Effective stack height  [m]"),
        ("z",  "Receptor height  [m]"),
    ]
    left = Inches(0.6)
    top  = Inches(5.7)
    for sym, desc in params:
        add_textbox(s, sym, left, top, Inches(0.7), Inches(0.35),
                    size=14, bold=True, color=BLUE)
        add_textbox(s, desc, left + Inches(0.75), top, Inches(5.0), Inches(0.35),
                    size=14, color=GRAY)
        top += Inches(0.38)

    params2 = [
        ("σ_y",  "Lateral spread  [m]"),
        ("σ_z",  "Vertical spread  [m]"),
        ("y",    "Crosswind offset  [m]"),
        ("x",    "Downwind distance  [m]"),
    ]
    left2 = Inches(6.8)
    top2  = Inches(5.7)
    for sym, desc in params2:
        add_textbox(s, sym, left2, top2, Inches(0.8), Inches(0.35),
                    size=14, bold=True, color=BLUE)
        add_textbox(s, desc, left2 + Inches(0.85), top2, Inches(5.0), Inches(0.35),
                    size=14, color=GRAY)
        top2 += Inches(0.38)


def slide_attenuation(prs):
    s = blank_slide(prs)
    fill_bg(s)
    section_bar(s, "6 · Optional Exponential Attenuations")

    # Chemical decay
    add_rect(s, Inches(0.3), Inches(1.05), Inches(12.7), Inches(0.01), ACCENT)
    add_textbox(s, "Chemical / radioactive decay   (first-order, rate  λ  [1/s])",
                Inches(0.4), Inches(1.1), Inches(12.0), Inches(0.45),
                size=17, bold=True, color=BLUE)

    add_latex(s,
        r"C' = C \cdot \exp\!\left(-\frac{\lambda\, x}{u}\right)",
        SLIDE_W//2, Inches(2.15), max_w=Inches(8), fontsize=26)

    add_textbox(s, "Travel time  τ = x/u  converts the spatial coordinate to time for first-order kinetics.",
                Inches(0.4), Inches(2.9), Inches(12.5), Inches(0.4),
                size=13, color=GRAY, italic=True)

    # Dry deposition
    add_rect(s, Inches(0.3), Inches(3.45), Inches(12.7), Inches(0.01), ACCENT)
    add_textbox(s, "Dry deposition   (velocity  v_d  [m/s],  mixing height  h_m  [m])",
                Inches(0.4), Inches(3.5), Inches(12.0), Inches(0.45),
                size=17, bold=True, color=BLUE)

    add_latex(s,
        r"C' = C \cdot \exp\!\left(-\frac{v_d\, x}{u\;\max(h_m,\,1)}\right)",
        SLIDE_W//2, Inches(4.55), max_w=Inches(9), fontsize=26)

    add_textbox(s, "Models vertically-uniform removal within the mixing layer (box approximation).",
                Inches(0.4), Inches(5.35), Inches(12.5), Inches(0.4),
                size=13, color=GRAY, italic=True)

    # Total
    add_rect(s, Inches(0.3), Inches(5.9), Inches(12.7), Inches(0.01), ACCENT)
    add_textbox(s, "Total concentration at receptor  (sum over all emittors)",
                Inches(0.4), Inches(5.95), Inches(12.0), Inches(0.45),
                size=17, bold=True, color=BLUE)

    add_latex(s,
        r"C_{total} = C_{bg} + \sum_{i}\, C_i(x_i, y_i, z_r)",
        SLIDE_W//2, Inches(6.7), max_w=Inches(9), fontsize=22)


def slide_assumptions(prs):
    s = blank_slide(prs)
    fill_bg(s)
    section_bar(s, "7 · Model Assumptions & Limitations")

    items = [
        ("Steady-state",
         "No time-of-flight: emissions are assumed to have reached steady state at each hour."),
        ("Homogeneous wind",
         "Speed and direction are uniform across the entire domain at each time step."),
        ("Flat terrain",
         "No topography, buildings, or obstacles deflect or shelter the plume."),
        ("Perfect ground reflection",
         "Image source at –H enforces zero vertical flux at z = 0 (no absorption)."),
        ("Prescribed plume rise",
         "ΔH is user-entered; Briggs' momentum / buoyancy rise is not computed."),
        ("Averaged wind speed for σ",
         "Stability class uses the mean wind speed across all emittors."),
        ("P-G coefficients",
         "Generic open-terrain Pasquill–Gifford values; not the Polish regulatory set."),
    ]

    top = Inches(1.1)
    for title, desc in items:
        add_rect(s, Inches(0.3), top, Inches(0.07), Inches(0.38), BLUE)
        add_textbox(s, title, Inches(0.5), top, Inches(3.2), Inches(0.38),
                    size=14, bold=True, color=YELLOW)
        add_textbox(s, desc, Inches(3.75), top, Inches(9.2), Inches(0.38),
                    size=14, color=GRAY)
        top += Inches(0.72)


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════

def main():
    prs = new_prs()

    slide_title(prs)
    slide_coordinate(prs)
    slide_windspeed(prs)
    slide_stability(prs)
    slide_sigma(prs)
    slide_concentration(prs)
    slide_attenuation(prs)
    slide_assumptions(prs)

    out = os.path.join(os.path.dirname(__file__),
                       "CitySim_GaussianPlume_Model.pptx")
    prs.save(out)
    print(f"Saved: {out}")

if __name__ == "__main__":
    main()
